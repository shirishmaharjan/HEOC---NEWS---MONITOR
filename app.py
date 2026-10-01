"""
HEOC News Monitor v3
Run UI:  streamlit run app.py          (administrators add ?admin=1 to the link)
Run CLI: python app.py --cli           (for scheduled runs)
Packages: streamlit requests beautifulsoup4 feedparser pandas trafilatura openpyxl
"""
import calendar
import html
import io
import json
import re
import sys
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from concurrent.futures import TimeoutError as FTimeout
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import quote, urljoin, urlparse

import feedparser
import pandas as pd
import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

try:
    import trafilatura
except ImportError:
    trafilatura = None

NPT = timezone(timedelta(hours=5, minutes=45))  # Nepal Time
BASE = Path(__file__).parent
CONFIG_FILE, SEEN_FILE, RUNS_DIR = BASE / "heoc_config.json", BASE / "seen_urls.txt", BASE / "runs"
TIERS = ["Critical", "Watch", "Routine"]
COLOR = {"Critical": "#d64545", "Watch": "#f0932b", "Routine": "#2ea36b"}
ICON = {"Critical": "🔴 Critical", "Watch": "🟠 Watch", "Routine": "🟢 Routine"}
META_COLS = ["priority", "source", "category", "title", "url", "published_npt", "date_note", "matched_keywords", "fetched_at", "file"]
SCOLS = ["source", "category", "status", "via", "checked", "matches", "note"]
COLS = ["enabled", "name", "category", "type", "url", "link_contains", "fulltext"]

# One shared session: browser-like headers + automatic retries (handles temporary 503 / timeouts)
SESSION = requests.Session()
SESSION.headers.update({
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36 HEOC-NewsMonitor/3.0",
    "Accept-Language": "en-US,en;q=0.8,ne;q=0.6"})
_retry = Retry(total=2, backoff_factor=1.0, status_forcelist=[429, 500, 502, 503, 504], allowed_methods=["GET"])
SESSION.mount("https://", HTTPAdapter(max_retries=_retry))
SESSION.mount("http://", HTTPAdapter(max_retries=_retry))


def now():
    return datetime.now(NPT)


def S(name, url, typ="auto", cat="National", contains="", full=True, on=True):
    return dict(enabled=on, name=name, url=url, type=typ, category=cat, link_contains=contains, fulltext=full)


def gnews(query, lang="en"):
    loc = "hl=en-NP&gl=NP&ceid=NP:en" if lang == "en" else "hl=ne&gl=NP&ceid=NP:ne"
    return f"https://news.google.com/rss/search?q={quote(query + ' when:7d')}&{loc}"


# type: rss = feed link | html = ordinary web page | auto = app finds the feed itself (best for new sites)
DEFAULTS = {
    "sources": [
        S("Kathmandu Post - Health", "https://kathmandupost.com/health", "html", contains="/health/"),
        S("Kathmandu Post - Feed", "https://kathmandupost.com/rss", "rss"),
        S("Onlinekhabar English", "https://english.onlinekhabar.com/feed", "rss"),
        S("Onlinekhabar (Nepali)", "https://www.onlinekhabar.com/feed", "rss"),
        S("Setopati (Nepali)", "https://www.setopati.com/feed", "rss"),
        S("Ratopati (Nepali)", "https://www.ratopati.com/feed", "rss"),
        S("Khabarhub", "https://khabarhub.com/feed", "rss"),
        S("Himalayan Times", "https://thehimalayantimes.com/rssFeed/15", "rss"),
        S("Republica", "https://myrepublica.nagariknetwork.com"),
        S("eKantipur", "https://ekantipur.com"),
        S("BBC Nepali", "https://feeds.bbci.co.uk/nepali/rss.xml", "rss"),
        S("Gorkhapatra", "https://gorkhapatraonline.com", "html"),
        S("Annapurna Post", "https://annapurnapost.com", "html"),
        S("Annapurna Express", "https://theannapurnaexpress.com"),
        S("Nagarik News", "https://nagariknews.nagariknetwork.com/feed", "rss"),
        S("Naya Patrika", "https://nayapatrikadaily.com"),
        S("Lokaantar", "https://lokaantar.com/feed", "rss"),
        S("Nepal Press", "https://nepalpress.com"),
        S("Nepali Times", "https://nepalitimes.com/feed", "rss"),
        S("Pahilo Post", "https://pahilopost.com"),
        S("Rajdhani Daily", "https://rajdhani.com.np/feed", "rss"),
        S("News of Nepal", "https://newsofnepal.com/feed", "rss"),
        S("The Rising Nepal", "https://risingnepaldaily.com", "html"),
        S("Rastriya Samachar Samiti (RSS)", "https://rss.com.np"),
        S("HEOC / MoHP", "https://heoc.mohp.gov.np", "html", "Government"),
        S("MoHP", "https://mohp.gov.np", "html", "Government"),
        S("EDCD", "https://edcd.gov.np", "html", "Government"),
        S("DoHS (Dept of Health Services)", "https://dohs.gov.np", "html", "Government"),
        S("DHM (Weather and Flood)", "https://www.dhm.gov.np", "html", "Government"),
        S("NHRC (Health Research Council)", "https://nhrc.gov.np", "html", "Government", on=False),
        S("NDRRMA (Disaster Authority)", "https://ndrrma.gov.np", "html", "Government", on=False),
        S("BIPAD Portal", "https://bipad.gov.np", "html", "Government", on=False),
        S("Google News - English", gnews("Nepal outbreak OR dengue OR cholera OR epidemic"), "rss", "Aggregator", full=False),
        S("Google News - Nepali", gnews("डेंगु OR हैजा OR प्रकोप OR महामारी", "ne"), "rss", "Aggregator", full=False),
        *[S(f"Google News - {p} Province", gnews(f"{p} Province health OR outbreak OR disease"), "rss", "Local", full=False)
          for p in ("Koshi", "Madhesh", "Bagmati", "Gandaki", "Lumbini", "Karnali", "Sudurpashchim")],
        S("Example local portal (edit me)", "https://example.com", "auto", "Local", on=False),
    ],
    "keywords": {
        "Critical": ["outbreak", "epidemic", "cholera", "dengue", "food poisoning", "प्रकोप", "महामारी", "हैजा", "डेंगु", "डेंगू", "खाद्य विषाक्तता"],
        "Watch": ["typhoid", "diarrhoea", "influenza", "measles", "rabies", "covid", "scrub typhus", "टाइफाइड", "झाडापखाला", "इन्फ्लुएन्जा", "दादुरा", "रेबिज", "कोरोना", "कोभिड"],
        "Routine": ["flood", "landslide", "earthquake", "बाढी", "पहिरो", "भूकम्प"],
    },
    "window_hours": 24, "fetch_full_text": True, "max_per_source": 40, "delay_sec": 1.0,
    "skip_seen": False, "include_undated": True, "workers": 4, "time_limit_sec": 600, "version": 4,
}


# ---------- config ----------
def load_config():
    cfg = json.loads(json.dumps(DEFAULTS))
    if CONFIG_FILE.exists():
        try:
            saved = json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
            ver = saved.get("version", 1)
            if isinstance(saved.get("keywords"), list):
                saved["keywords"] = {"Critical": [], "Watch": saved["keywords"], "Routine": []}
            if "days_back" in saved and "window_hours" not in saved:
                saved["window_hours"] = int(saved["days_back"]) * 24
            cfg.update({k: v for k, v in saved.items() if k in cfg and k != "version"})
            if ver < DEFAULTS["version"]:  # older config: refresh built-in sources, keep the user's own rows and On/Off choices
                names = {d["name"] for d in DEFAULTS["sources"]}
                old = {s.get("name"): s for s in cfg["sources"]}
                custom = [s for s in cfg["sources"] if s.get("name") not in names]
                fresh = []
                for d in DEFAULTS["sources"]:
                    d = dict(d)
                    d["enabled"] = old.get(d["name"], {}).get("enabled", d["enabled"]) if d["enabled"] else False
                    fresh.append(d)
                cfg["sources"], cfg["skip_seen"] = fresh + custom, False
        except Exception:
            pass
    if not any(cfg["keywords"].get(t) for t in TIERS):
        cfg["keywords"] = json.loads(json.dumps(DEFAULTS["keywords"]))
    cfg["sources"] = [{**S("", ""), **s} for s in cfg["sources"]]
    return cfg


def save_config(cfg):
    CONFIG_FILE.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")


# ---------- fetching ----------
def short(ex):
    return f"{type(ex).__name__}: {str(ex)[:110]}"


def explain(note):
    n = (note or "").lower()
    if "404" in n or "no entries" in n or "valid rss" in n:
        return "link has changed"
    if any(x in n for x in ("503", "429", "timeout", "retry", "connection")):
        return "site busy or blocking - try later"
    if "javascript" in n:
        return "site needs JavaScript"
    return "temporarily unavailable"


def to_npt(dt):
    if dt is None:
        return None
    return dt.replace(tzinfo=NPT) if dt.tzinfo is None else dt.astimezone(NPT)


def get(url, timeout=(10, 25)):
    r = SESSION.get(url, timeout=timeout)
    r.raise_for_status()
    if not r.encoding or r.encoding.lower() == "iso-8859-1":
        r.encoding = r.apparent_encoding or "utf-8"  # keeps Nepali text readable
    return r


def fetch_rss(url):
    feed = feedparser.parse(get(url).content)
    if not feed.entries and not feed.get("version"):
        raise ValueError("not a valid RSS/Atom feed")
    items = []
    for e in feed.entries:
        t = e.get("published_parsed") or e.get("updated_parsed")
        pub = datetime.fromtimestamp(calendar.timegm(t), NPT) if t else None
        summary = BeautifulSoup(e.get("summary", ""), "html.parser").get_text(" ")
        items.append({"title": e.get("title", ""), "url": e.get("link", ""), "published": pub, "summary": summary})
    return items


SKIP_LINK = re.compile(r"(login|signin|register|/tag/|/category/|/author/|/page/\d|javascript:|mailto:|\.(jpg|jpeg|png|gif|mp4)$)", re.I)


def fetch_html_listing(src):
    soup = BeautifulSoup(get(src["url"]).text, "html.parser")
    domain, must, seen, items = urlparse(src["url"]).netloc, src.get("link_contains", ""), set(), []
    for a in soup.find_all("a", href=True):
        link = urljoin(src["url"], a["href"]).split("#")[0]
        title = a.get_text(" ", strip=True)
        if urlparse(link).netloc != domain or link in seen or len(title) < 20 or SKIP_LINK.search(link) or (must and must not in link):
            continue
        seen.add(link)
        items.append({"title": title, "url": link, "published": None, "summary": ""})
    if not items:
        raise ValueError("No article links found (site may need JavaScript)")
    return items


def discover_feed(url):
    """Look inside a web page for its official RSS/Atom link."""
    soup = BeautifulSoup(get(url).text, "html.parser")
    for l in soup.find_all("link"):
        typ, href = (l.get("type") or "").lower(), l.get("href")
        if href and ("rss" in typ or "atom" in typ) and "comment" not in (href + (l.get("title") or "")).lower():
            return urljoin(url, href)
    return None


def fetch_items(src):
    """Try the chosen method first, then fall back to the others, so one wrong link does not lose the source."""
    url, typ = src["url"], src.get("type", "auto")
    home = f"{urlparse(url).scheme}://{urlparse(url).netloc}/"
    order = {"rss": ["rss", "discover", "html"], "html": ["html", "discover"]}.get(typ, ["discover", "html"])
    errs = []
    for step in order:
        try:
            if step == "rss":
                return fetch_rss(url), "RSS"
            if step == "html":
                return fetch_html_listing({**src, "url": url if typ != "rss" else home}), "Web page"
            for page in dict.fromkeys([url, home]):
                try:
                    feed = discover_feed(page)
                except Exception:
                    continue
                if feed and feed != url:
                    return fetch_rss(feed), "RSS (auto-found)"
            raise ValueError("no feed found")
        except Exception as ex:
            errs.append(short(ex))
    raise ValueError(" | ".join(errs)[:220])


def fetch_article(url):
    html_text = get(url).text
    soup = BeautifulSoup(html_text, "html.parser")
    pub = None
    for tag in (soup.find("meta", attrs={"property": "article:published_time"}),
                soup.find("meta", attrs={"name": "article:published_time"}),
                soup.find("meta", attrs={"itemprop": "datePublished"}),
                soup.find("time", attrs={"datetime": True})):
        if tag is None:
            continue
        try:
            pub = to_npt(datetime.fromisoformat((tag.get("content") or tag.get("datetime")).strip().replace("Z", "+00:00")))
            break
        except Exception:
            continue
    text = trafilatura.extract(html_text) if trafilatura else None
    if not text:
        text = "\n".join(p.get_text(" ", strip=True) for p in soup.find_all("p"))
    return text or "", pub


# ---------- scanning ----------
def find_matches(text, kw):
    t = text.lower()
    for tier in TIERS:  # highest tier that matches decides the priority
        if any(k.lower() in t for k in kw[tier]):
            return tier, [k for tt in TIERS for k in kw[tt] if k.lower() in t]
    return None, []


def scan_source(src, cfg, cutoff, seen, kw):
    st = dict(source=src["name"], category=src.get("category", ""), status="OK", via="", checked=0, matches=0, note="")
    hits = []
    try:
        items, st["via"] = fetch_items(src)
        items = items[: int(cfg["max_per_source"])]
        st["checked"] = len(items)
        if not items:
            st["note"] = "no items in feed right now"
        for it in items:
            url, pub = it["url"], it["published"]
            if not url or url in seen or (pub and pub < cutoff):
                continue
            head, text = it["title"] + " " + it["summary"], ""
            if cfg["fetch_full_text"] and src.get("fulltext", True) and not url.lower().endswith(".pdf"):
                try:
                    text, art_pub = fetch_article(url)
                    pub = pub or art_pub
                    time.sleep(cfg["delay_sec"])
                except Exception:
                    pass
                if pub and pub < cutoff:
                    continue
            if not pub and not cfg.get("include_undated", True):
                continue
            prio, matched = find_matches(head + " " + text, kw)
            if prio:
                hits.append(dict(priority=prio, source=src["name"], category=src.get("category", ""), title=it["title"],
                                 url=url, published=pub, text=text or it["summary"], matched=matched))
        st["matches"] = len(hits)
    except Exception as ex:
        st.update(status="FAILED", note=str(ex)[:220])
    return st, hits


def slug(s, n):
    return re.sub(r"[^\w]+", "_", s, flags=re.UNICODE).strip("_")[:n] or "x"


def write_outputs(out_dir, hits, statuses):
    """Write meta.csv, source_status.csv and article files. Safe to call repeatedly (keeps partial results on disk)."""
    uniq = {h["url"]: h for h in hits}.values()
    ordered = sorted(uniq, key=lambda h: (TIERS.index(h["priority"]), -(h["published"].timestamp() if h["published"] else 0)))
    for old in (out_dir / "articles").glob("*.txt"):
        old.unlink()
    rows = []
    for n, h in enumerate(ordered, 1):
        fname = f"{n:03d}_{slug(h['source'], 25)}_{slug(h['title'], 50)}.txt"
        pub = h["published"].strftime("%Y-%m-%d %H:%M") if h["published"] else ""
        (out_dir / "articles" / fname).write_text(
            f"Title: {h['title']}\nSource: {h['source']}\nURL: {h['url']}\nPublished (NPT): {pub}\n"
            f"Priority: {h['priority']}\nKeywords: {', '.join(h['matched'])}\n\n{h['text']}", encoding="utf-8")
        rows.append(dict(priority=h["priority"], source=h["source"], category=h["category"], title=h["title"], url=h["url"],
                         published_npt=pub, date_note="" if pub else "date not found - please verify",
                         matched_keywords=", ".join(h["matched"]), fetched_at=now().strftime("%Y-%m-%d %H:%M:%S"), file=fname))
    df, sdf = pd.DataFrame(rows, columns=META_COLS), pd.DataFrame(statuses, columns=SCOLS)
    df.to_csv(out_dir / "meta.csv", index=False, encoding="utf-8-sig")
    sdf.to_csv(out_dir / "source_status.csv", index=False, encoding="utf-8-sig")
    return df, sdf


def run_scan(cfg, progress=None):
    kw = {t: [k.strip() for k in cfg["keywords"].get(t, []) if k.strip()] for t in TIERS}
    cutoff = now() - timedelta(hours=float(cfg["window_hours"]))
    seen = set(SEEN_FILE.read_text(encoding="utf-8").split()) if (cfg["skip_seen"] and SEEN_FILE.exists()) else set()
    sources = [s for s in cfg["sources"] if s.get("enabled") and s.get("url")]
    out_dir = RUNS_DIR / now().strftime("%Y%m%d_%H%M%S")
    (out_dir / "articles").mkdir(parents=True, exist_ok=True)

    statuses, hits = [], []
    ex = ThreadPoolExecutor(max_workers=int(cfg["workers"]))
    futs = {ex.submit(scan_source, s, cfg, cutoff, seen, kw): s for s in sources}
    try:
        for i, f in enumerate(as_completed(futs, timeout=float(cfg["time_limit_sec"])), 1):
            try:
                st, h = f.result()
            except Exception as e:  # a single source can never stop the scan
                st, h = dict(source=futs[f]["name"], category="", status="FAILED", via="", checked=0, matches=0, note=short(e)), []
            statuses.append(st)
            hits += h
            if h:
                write_outputs(out_dir, hits, statuses)  # partial results stay on disk if the page disconnects
            if progress:
                progress(i / len(futs), st["source"])
    except FTimeout:  # time limit reached: keep what we have
        for f, s in futs.items():
            if not f.done():
                statuses.append(dict(source=s["name"], category=s.get("category", ""), status="TIMED OUT", via="",
                                     checked=0, matches=0, note="Skipped: scan time limit reached"))
    finally:
        ex.shutdown(wait=False, cancel_futures=True)

    df, sdf = write_outputs(out_dir, hits, statuses)
    if len(df):
        with SEEN_FILE.open("a", encoding="utf-8") as f:
            f.write("\n".join(df["url"]) + "\n")
    return df, sdf, out_dir


def test_sources(cfg):
    def one(s):
        try:
            items, via = fetch_items(s)
            return dict(source=s["name"], status="OK", via=via, items_found=len(items), sample=(items[0]["title"][:70] if items else ""))
        except Exception as ex:
            return dict(source=s["name"], status="FAILED", via="", items_found=0, sample=explain(str(ex)))
    srcs = [s for s in cfg["sources"] if s.get("enabled") and s.get("url")]
    with ThreadPoolExecutor(max_workers=6) as ex:
        return pd.DataFrame(list(ex.map(one, srcs)))


# ---------- exports ----------
def make_zip(out_dir):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for p in out_dir.rglob("*"):
            if p.is_file():
                z.write(p, p.relative_to(out_dir))
    return buf.getvalue()


def fmt_run(name):
    try:
        return datetime.strptime(name, "%Y%m%d_%H%M%S").strftime("%d %b %Y, %H:%M")
    except ValueError:
        return name


def file_stamp(name):
    """Date (and time) of the scan for file names, e.g. 2026-10-01_0830."""
    try:
        return datetime.strptime(name, "%Y%m%d_%H%M%S").strftime("%Y-%m-%d_%H%M")
    except ValueError:
        return now().strftime("%Y-%m-%d_%H%M")


def to_excel(df):
    try:
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return None
    cols = {"priority": "Priority", "source": "Source", "category": "Category", "title": "Title", "url": "Link",
            "published_npt": "Published (Nepal time)", "date_note": "Note", "matched_keywords": "Keywords matched"}
    out = df[list(cols)].rename(columns=cols)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        out.to_excel(w, index=False, sheet_name="HEOC News")
        ws = w.sheets["HEOC News"]
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0B3C5D")
        fills = {"Critical": "F8D7DA", "Watch": "FFE8CC", "Routine": "D6F0E0"}
        for row in ws.iter_rows(min_row=2):
            row[0].fill = PatternFill("solid", fgColor=fills.get(row[0].value, "FFFFFF"))
            if row[4].value:
                row[4].hyperlink = row[4].value
                row[4].font = Font(color="0563C1", underline="single")
            row[3].alignment = Alignment(wrap_text=True, vertical="top")
        for letter, width in zip("ABCDEFGH", (11, 26, 13, 70, 18, 20, 28, 30)):
            ws.column_dimensions[letter].width = width
        ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
    return buf.getvalue()


def latest_result():
    runs = sorted(RUNS_DIR.glob("*/meta.csv"), reverse=True) if RUNS_DIR.exists() else []
    for p in runs:
        try:
            df = pd.read_csv(p).fillna("")
            sp = p.parent / "source_status.csv"
            sdf = pd.read_csv(sp).fillna("") if sp.exists() else pd.DataFrame(columns=SCOLS)
            return df, sdf, p.parent.name, make_zip(p.parent)
        except Exception:
            continue
    return None


# ---------- UI ----------
CSS = """<style>
.block-container{padding-top:1.4rem;max-width:1250px}
footer{visibility:hidden}
.hdr{background:linear-gradient(100deg,#0b3c5d 0%,#1d6fa5 100%);padding:20px 26px;border-radius:14px;color:#fff;margin-bottom:14px;display:flex;justify-content:space-between;align-items:center;flex-wrap:wrap;gap:8px}
.hdr h2{margin:0;color:#fff;font-size:1.55rem}.hdr p{margin:3px 0 0;opacity:.88}.hdr .clock{text-align:right;font-size:.95rem;opacity:.95}
.kpi{border:1px solid rgba(128,128,128,.25);border-radius:12px;padding:14px 16px;background:rgba(128,128,128,.07)}
.kpi b{font-size:1.9rem;display:block;line-height:1.1}.kpi span{opacity:.75;font-size:.85rem}
.card{border:1px solid rgba(128,128,128,.22);border-left:6px solid #888;border-radius:10px;padding:11px 16px;margin:8px 0;background:rgba(128,128,128,.05)}
.card a{font-weight:600;font-size:1.02rem;text-decoration:none}.card a:hover{text-decoration:underline}
.badge{color:#fff;border-radius:20px;padding:2px 10px;font-size:.72rem;font-weight:700;margin-right:8px;vertical-align:middle}
.meta{opacity:.7;font-size:.82rem;margin:4px 0}
.chip{display:inline-block;background:rgba(29,111,165,.15);border-radius:12px;padding:1px 9px;font-size:.75rem;margin:2px 4px 0 0}
</style>"""


def kpi(col, value, label, color=None):
    style = f' style="color:{color}"' if color else ""
    col.markdown(f'<div class="kpi"><b{style}>{value}</b><span>{label}</span></div>', unsafe_allow_html=True)


def card_html(r):
    c = COLOR.get(r["priority"], "#888")
    e = lambda x: html.escape(str(x), quote=True)
    chips = "".join(f'<span class="chip">{e(k.strip())}</span>' for k in str(r["matched_keywords"]).split(",") if k.strip())
    when = r["published_npt"] or "date not found - please verify"
    return (f'<div class="card" style="border-left-color:{c}"><span class="badge" style="background:{c}">{e(r["priority"])}</span>'
            f'<a href="{e(r["url"])}" target="_blank" rel="noopener">{e(r["title"])}</a>'
            f'<div class="meta">📰 {e(r["source"])} · {e(r["category"])} · 🕒 {e(when)}</div><div>{chips}</div></div>')


def render_results(df, sdf, name, zbytes, key):
    import streamlit as st

    st.caption(f"📅 Scan from **{fmt_run(name)}** (Nepal time)")
    ok, bad = (sdf.status == "OK").sum() if len(sdf) else 0, sdf[sdf.status != "OK"] if len(sdf) else sdf
    crit = int((df.priority == "Critical").sum()) if len(df) else 0
    k = st.columns(4)
    kpi(k[0], len(df), "Articles found")
    kpi(k[1], crit, "Critical alerts", COLOR["Critical"] if crit else None)
    kpi(k[2], int(ok), "Sources read")
    kpi(k[3], len(bad), "Sources unavailable", COLOR["Watch"] if len(bad) else None)
    if len(bad):
        names = ", ".join(f"{r.source} ({explain(r.note)})" for r in bad.itertuples())
        st.warning(f"⚠️ {len(bad)} source(s) could not be read this time: {names}. The scan continued with all the others.")
    st.write("")
    if not len(df):
        st.success("✅ No matching articles in this time window.")
    else:
        f1, f2, f3 = st.columns([2, 2, 1.4])
        pick = f1.multiselect("Priority", TIERS, default=TIERS, key=f"{key}_p")
        q = f2.text_input("🔎 Search in results", key=f"{key}_q", placeholder="e.g. dengue, Lalitpur, डेंगु")
        view_mode = f3.radio("View", ["Cards", "Table"], horizontal=True, key=f"{key}_v")
        v = df[df.priority.isin(pick)]
        if q.strip():
            blob = v.title + " " + v.matched_keywords + " " + v.source
            v = v[blob.str.contains(q.strip(), case=False, na=False, regex=False)]
        st.caption(f"Showing {len(v)} of {len(df)} articles")
        if view_mode == "Cards":
            st.markdown("".join(card_html(r) for r in v.head(150).to_dict("records")), unsafe_allow_html=True)
            if len(v) > 150:
                st.info("Showing the first 150. Use Table view or download the file to see all.")
        else:
            t = v.drop(columns=["file", "fetched_at"], errors="ignore").copy()
            t["priority"] = t["priority"].map(ICON)
            st.dataframe(t, hide_index=True, column_config={
                "priority": "Priority", "source": "Source", "category": "Category", "title": "Title",
                "url": st.column_config.LinkColumn("Link", display_text="Open"), "published_npt": "Published (NPT)",
                "date_note": "Note", "matched_keywords": "Keywords"})
        stamp = file_stamp(name)
        d1, d2, d3 = st.columns(3)
        xl = to_excel(df)
        if xl:
            d1.download_button("⬇️ Excel (colour-coded)", xl, f"HEOC_News_{stamp}.xlsx",
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key=f"{key}_x")
        d2.download_button("⬇️ CSV", df.to_csv(index=False, encoding="utf-8-sig"), f"HEOC_News_{stamp}.csv", "text/csv", key=f"{key}_c")
        d3.download_button("⬇️ ZIP (CSV + full articles)", zbytes, f"HEOC_News_{stamp}.zip", "application/zip", key=f"{key}_z")
    with st.expander("Source status (which sites worked)"):
        st.dataframe(sdf, hide_index=True)


def main_ui():
    import streamlit as st

    st.set_page_config(page_title="HEOC News Monitor", page_icon="🩺", layout="wide")
    try:  # optional password: set APP_PASSWORD in Streamlit secrets
        pw = st.secrets.get("APP_PASSWORD")
    except Exception:
        pw = None
    if pw and st.text_input("Password", type="password") != pw:
        st.stop()

    @st.cache_resource
    def get_lock():
        import threading
        return threading.Lock()

    st.markdown(CSS, unsafe_allow_html=True)
    n = now()
    st.markdown(
        '<div class="hdr"><div><h2>🩺 Health Emergency Operation Centre - News Monitor</h2>'
        '<p>Automatic scan of national and local news for health alerts</p></div>'
        f'<div class="clock">{n.strftime("%A, %d %B %Y")}<br>🕒 {n.strftime("%H:%M")} Nepal time</div></div>', unsafe_allow_html=True)

    if "cfg" not in st.session_state:
        st.session_state.cfg = load_config()
        st.session_state.src_df = pd.DataFrame(st.session_state.cfg["sources"])[COLS]
    cfg = st.session_state.cfg
    admin = st.query_params.get("admin") == "1"  # administrators open the link with ?admin=1 at the end

    if admin:
        with st.sidebar:
            st.header("🔧 Administrator")
            cfg["max_per_source"] = st.number_input("Max articles per source", 5, 200, int(cfg["max_per_source"]))
            cfg["delay_sec"] = st.slider("Delay between requests (sec)", 0.0, 5.0, float(cfg["delay_sec"]), 0.5)
            cfg["workers"] = st.slider("Sources scanned in parallel", 1, 8, int(cfg["workers"]))
            cfg["time_limit_sec"] = st.number_input("Scan time limit (seconds)", 60, 3600, int(cfg["time_limit_sec"]), 60)
            if st.button("💾 Save settings"):
                save_config(cfg)
                st.success("Saved")
            if st.button("↩️ Reset to defaults"):
                if CONFIG_FILE.exists():
                    CONFIG_FILE.unlink()
                for k in list(st.session_state.keys()):
                    if k in ("cfg", "src_df", "result", "src_editor") or k.startswith("kw_"):
                        st.session_state.pop(k, None)
                st.rerun()

    tabs = st.tabs(["🏠 Dashboard", "🗂 History", "❓ Help"] + (["📰 Sources", "🔑 Keywords"] if admin else []))
    t_dash, t_hist, t_help = tabs[:3]

    if admin:
        t_src, t_kw = tabs[3:]
        with t_src:
            st.caption("Tick **On** to include a source. Type: `auto` = app finds the feed itself (recommended for new sites), `rss` = feed link, `html` = ordinary page.")
            edited = st.data_editor(
                st.session_state.src_df, num_rows="dynamic", key="src_editor",
                column_config={
                    "enabled": st.column_config.CheckboxColumn("On", default=True),
                    "name": st.column_config.TextColumn("Name", required=True),
                    "category": st.column_config.SelectboxColumn("Category", options=["National", "Local", "Government", "Aggregator"], default="Local"),
                    "type": st.column_config.SelectboxColumn("Type", options=["auto", "rss", "html"], default="auto", required=True),
                    "url": st.column_config.TextColumn("Website / feed link", required=True),
                    "link_contains": st.column_config.TextColumn("Link must contain"),
                    "fulltext": st.column_config.CheckboxColumn("Read full text", default=True)})
            e = edited.fillna({"enabled": True, "fulltext": True, "type": "auto", "category": "Local", "name": "", "url": "", "link_contains": ""})
            cfg["sources"] = [r for r in e.to_dict("records") if str(r["url"]).strip()]
            if st.button("🧪 Test all sources"):
                with st.spinner("Checking each source..."):
                    st.dataframe(test_sources(cfg), hide_index=True)
        with t_kw:
            st.caption("One keyword per line, English or Nepali. An article's priority is the highest level it matches.")

            def _restore_kw():
                for t in TIERS:
                    st.session_state[f"kw_{t}"] = "\n".join(DEFAULTS["keywords"][t])

            st.button("↩️ Restore default keywords", on_click=_restore_kw)
            cols, kw = st.columns(3), {}
            for c, tier, hint in zip(cols, TIERS, ["Act first", "Keep an eye on", "Background"]):
                with c:
                    txt = st.text_area(f"{ICON[tier]} ({hint})", "\n".join(cfg["keywords"].get(tier, [])), height=320, key=f"kw_{tier}")
                    kw[tier] = [x.strip() for x in txt.split("\n") if x.strip()]
            cfg["keywords"] = kw

    with t_dash:
        presets = {"Last 2 hours": 2, "Last 6 hours": 6, "Last 12 hours": 12, "Last 24 hours": 24,
                   "Last 48 hours": 48, "Last 7 days": 168, "Custom (hours)": None}
        with st.container(border=True):
            c1, c2, c3, c4 = st.columns([2, 2.2, 2.6, 1.3])
            label = c1.selectbox("📅 Time window", list(presets), index=3)
            hours = presets[label] or c1.number_input("Hours", 1, 720, 24)
            mode = c2.radio("Scan depth", ["Thorough (reads full articles)", "Quick (headlines only)"], index=0 if cfg["fetch_full_text"] else 1)
            cfg["window_hours"], cfg["fetch_full_text"] = hours, mode.startswith("Thorough")
            cfg["skip_seen"] = c3.checkbox("Only show articles not seen in earlier scans", cfg["skip_seen"])
            cfg["include_undated"] = c3.checkbox("Include articles with no publish date", cfg["include_undated"])
            c4.write("")
            run = c4.button("▶ Run scan", type="primary")

        if run:
            lock = get_lock()
            if not lock.acquire(blocking=False):
                st.warning("⏳ Another scan is running right now. Please wait a minute and press Run scan again.")
            else:
                bar = st.progress(0.0, text="Starting scan...")
                try:
                    df, sdf, out = run_scan(cfg, lambda p, s: bar.progress(p, text=f"Checked: {s}"))
                    st.session_state.result = (df, sdf, out.name, make_zip(out))
                except Exception as ex:
                    st.error("The scan could not be completed. Please try again. If it keeps happening, tell the administrator.")
                    with st.expander("Technical details"):
                        st.code(short(ex))
                finally:
                    bar.empty()
                    lock.release()

        if "result" not in st.session_state:
            lr = latest_result()
            if lr:
                st.session_state.result = lr
        if "result" in st.session_state:
            render_results(*st.session_state.result, key="d")
        else:
            st.info("No scan yet. Choose a time window and press **▶ Run scan**.")

    with t_hist:
        runs = sorted(RUNS_DIR.glob("*/meta.csv"), reverse=True) if RUNS_DIR.exists() else []
        if not runs:
            st.info("No saved scans yet. (On Streamlit Cloud, history resets when the app restarts.)")
        else:
            sel = st.selectbox("Choose a scan", [p.parent.name for p in runs], format_func=fmt_run)
            try:
                folder = RUNS_DIR / sel
                h = pd.read_csv(folder / "meta.csv").fillna("")
                sp = folder / "source_status.csv"
                s = pd.read_csv(sp).fillna("") if sp.exists() else pd.DataFrame(columns=SCOLS)
                render_results(h, s, sel, make_zip(folder), key="h")
            except Exception as ex:
                st.warning(f"Could not read this scan: {short(ex)}")

    with t_help:
        st.markdown("""
**How to use**
1. Choose the **time window** (e.g. last 2 hours) and press **▶ Run scan**.
2. Review the results. 🔴 Critical items come first. Click the title to read the news.
3. Download **Excel**, **CSV** or **ZIP**. File names include the scan date, e.g. `HEOC_News_2026-10-01_0830.xlsx`.

**What the colours mean**
- 🔴 **Critical**: outbreak words such as dengue, cholera, epidemic. Check first.
- 🟠 **Watch**: other diseases to keep an eye on.
- 🟢 **Routine**: background events such as floods or landslides.

**Good to know**
- "Quick" is faster but only reads headlines. "Thorough" reads full articles and finds more.
- Articles without a publish date are marked, so please verify them.
- A source shown as unavailable is usually a temporary website problem. The rest still work.
- This tool is a first filter. A researcher should confirm each item before reporting.
""")


if __name__ == "__main__":
    if "--cli" in sys.argv:
        d, s, o = run_scan(load_config())
        print(f"Done: {len(d)} articles, {int((s.status != 'OK').sum())} unavailable sources. Saved in {o}")
    else:
        main_ui()
